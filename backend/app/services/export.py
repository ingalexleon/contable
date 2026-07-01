import io
from typing import List, Dict, Any

from openpyxl import Workbook
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select

from app.models.client import Client
from app.models.payment import Payment
from app.models.service import Service, ClientService


async def export_clients_to_excel(db: AsyncSession) -> io.BytesIO:
    """Export all active clients to an Excel file."""
    result = await db.execute(select(Client).where(Client.is_active == True))
    clients = result.scalars().all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Clientes"

    headers = [
        "ID", "Razon Social", "Contacto", "Email",
        "Telefono", "RFC", "Tipo", "Direccion", "Fecha Registro"
    ]
    ws.append(headers)

    for client in clients:
        ws.append([
            client.id,
            client.business_name,
            client.contact_name,
            client.email,
            client.phone,
            client.rfc,
            client.client_type.value if client.client_type else "",
            client.address,
            client.created_at.strftime("%Y-%m-%d") if client.created_at else "",
        ])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


async def export_payments_to_excel(db: AsyncSession) -> io.BytesIO:
    """Export all payments to an Excel file."""
    result = await db.execute(select(Payment))
    payments = result.scalars().all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Pagos"

    headers = [
        "ID", "Cliente ID", "Monto", "Fecha Pago",
        "Periodo Inicio", "Periodo Fin", "Estado", "Notas"
    ]
    ws.append(headers)

    for payment in payments:
        ws.append([
            payment.id,
            payment.client_id,
            payment.amount,
            str(payment.payment_date) if payment.payment_date else "",
            str(payment.period_start) if payment.period_start else "",
            str(payment.period_end) if payment.period_end else "",
            payment.status.value if payment.status else "",
            payment.notes,
        ])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


async def export_services_to_excel(db: AsyncSession) -> io.BytesIO:
    """Export all services and client-service assignments to an Excel file."""
    # Services catalog
    services_result = await db.execute(select(Service).where(Service.is_active == True))
    services = services_result.scalars().all()

    # Client-service assignments
    cs_result = await db.execute(
        select(ClientService, Service, Client)
        .join(Service, ClientService.service_id == Service.id)
        .join(Client, ClientService.client_id == Client.id)
        .where(ClientService.is_active == True)
    )
    assignments = cs_result.all()

    wb = Workbook()

    # Services catalog sheet
    ws_services = wb.active
    ws_services.title = "Catalogo Servicios"
    ws_services.append(["ID", "Nombre", "Descripcion", "Precio Base", "Categoria"])
    for svc in services:
        ws_services.append([
            svc.id,
            svc.name,
            svc.description or "",
            svc.default_price,
            svc.category or "",
        ])

    # Client-service assignments sheet
    ws_assignments = wb.create_sheet("Asignaciones")
    ws_assignments.append([
        "ID", "Cliente", "Servicio", "Precio Personalizado",
        "Periodo", "Fecha Inicio"
    ])
    for cs, service, client in assignments:
        ws_assignments.append([
            cs.id,
            client.business_name,
            service.name,
            cs.custom_price if cs.custom_price is not None else service.default_price,
            cs.period.value if cs.period else "",
            str(cs.start_date) if cs.start_date else "",
        ])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


async def export_all_to_excel(db: AsyncSession) -> io.BytesIO:
    """Export all data (clients, payments, services) in a multi-sheet Excel file."""
    # Clients
    clients_result = await db.execute(select(Client).where(Client.is_active == True))
    clients = clients_result.scalars().all()

    # Payments
    payments_result = await db.execute(select(Payment))
    payments = payments_result.scalars().all()

    # Services
    services_result = await db.execute(select(Service).where(Service.is_active == True))
    services = services_result.scalars().all()

    # Client-service assignments
    cs_result = await db.execute(
        select(ClientService, Service, Client)
        .join(Service, ClientService.service_id == Service.id)
        .join(Client, ClientService.client_id == Client.id)
        .where(ClientService.is_active == True)
    )
    assignments = cs_result.all()

    wb = Workbook()

    # Clients sheet
    ws_clients = wb.active
    ws_clients.title = "Clientes"
    ws_clients.append([
        "ID", "Razon Social", "Contacto", "Email",
        "Telefono", "RFC", "Tipo", "Direccion", "Fecha Registro"
    ])
    for client in clients:
        ws_clients.append([
            client.id,
            client.business_name,
            client.contact_name,
            client.email,
            client.phone,
            client.rfc,
            client.client_type.value if client.client_type else "",
            client.address,
            client.created_at.strftime("%Y-%m-%d") if client.created_at else "",
        ])

    # Payments sheet
    ws_payments = wb.create_sheet("Pagos")
    ws_payments.append([
        "ID", "Cliente ID", "Monto", "Fecha Pago",
        "Periodo Inicio", "Periodo Fin", "Estado", "Notas"
    ])
    for payment in payments:
        ws_payments.append([
            payment.id,
            payment.client_id,
            payment.amount,
            str(payment.payment_date) if payment.payment_date else "",
            str(payment.period_start) if payment.period_start else "",
            str(payment.period_end) if payment.period_end else "",
            payment.status.value if payment.status else "",
            payment.notes,
        ])

    # Services sheet
    ws_services = wb.create_sheet("Servicios")
    ws_services.append(["ID", "Nombre", "Descripcion", "Precio Base", "Categoria"])
    for svc in services:
        ws_services.append([
            svc.id,
            svc.name,
            svc.description or "",
            svc.default_price,
            svc.category or "",
        ])

    # Assignments sheet
    ws_assignments = wb.create_sheet("Asignaciones")
    ws_assignments.append([
        "ID", "Cliente", "Servicio", "Precio Personalizado",
        "Periodo", "Fecha Inicio"
    ])
    for cs, service, client in assignments:
        ws_assignments.append([
            cs.id,
            client.business_name,
            service.name,
            cs.custom_price if cs.custom_price is not None else service.default_price,
            cs.period.value if cs.period else "",
            str(cs.start_date) if cs.start_date else "",
        ])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


async def export_general_report_excel(
    db: AsyncSession,
    clients_data: List[Dict[str, Any]],
    payments_data: List[Dict[str, Any]],
) -> io.BytesIO:
    """Export a general report with clients and payments sheets."""
    wb = Workbook()

    # Clients sheet
    ws_clients = wb.active
    ws_clients.title = "Clientes"
    if clients_data:
        ws_clients.append(list(clients_data[0].keys()))
        for row in clients_data:
            ws_clients.append(list(row.values()))

    # Payments sheet
    ws_payments = wb.create_sheet("Pagos")
    if payments_data:
        ws_payments.append(list(payments_data[0].keys()))
        for row in payments_data:
            ws_payments.append(list(row.values()))

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
